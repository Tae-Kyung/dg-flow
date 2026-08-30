import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

doc = Document()

# ── 페이지 설정 (A4 가로) ──
for section in doc.sections:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# ── 스타일 설정 ──
style = doc.styles['Normal']
font = style.font
font.name = '맑은 고딕'
font.size = Pt(10)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def set_cell_shading(cell, color):
    """셀 배경색 설정"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_border(cell, **kwargs):
    """셀 테두리 설정"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, val in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{val.get("val", "single")}" '
            f'w:sz="{val.get("sz", "4")}" w:space="0" '
            f'w:color="{val.get("color", "000000")}"/>'
        )
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_paragraph(text, size=10, bold=False, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0, color=None):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = Pt(size * 1.6)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    if color:
        run.font.color.rgb = color
    return p

def add_run_to_paragraph(p, text, size=10, bold=False, color=None):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    if color:
        run.font.color.rgb = color
    return run

def format_cell(cell, text, size=10, bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    if color:
        run.font.color.rgb = color
    return cell

# ===================================================
# 헤더 영역
# ===================================================

p = add_paragraph("지역혁신중심 대학지원체계(RISE) AX기업혁신 공동연구센터",
                   size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

# 제목 박스 (테이블로 구현)
title_table = doc.add_table(rows=1, cols=1)
title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = title_table.cell(0, 0)
set_cell_shading(cell, "2B2B2B")
format_cell(cell, "「동일유리 업무자동화(AX)」 세미나 운영 계획(안)",
            size=16, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
cell.paragraphs[0].paragraph_format.space_before = Pt(6)
cell.paragraphs[0].paragraph_format.space_after = Pt(6)

# 날짜
add_paragraph("2026. 08. 29(토)", size=11, alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_before=8, space_after=12)

# ===================================================
# 1. 추진배경 및 목적
# ===================================================

# 섹션 번호 테이블
def add_section_header(number, title):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    # 번호 셀
    c1 = t.cell(0, 0)
    set_cell_shading(c1, "4472C4")
    format_cell(c1, str(number), size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    c1.width = Cm(1.2)
    c1.paragraphs[0].paragraph_format.space_before = Pt(4)
    c1.paragraphs[0].paragraph_format.space_after = Pt(4)
    # 제목 셀
    c2 = t.cell(0, 1)
    format_cell(c2, f"  {title}", size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    c2.paragraphs[0].paragraph_format.space_before = Pt(4)
    c2.paragraphs[0].paragraph_format.space_after = Pt(4)
    # 테이블 너비
    t.columns[0].width = Cm(1.2)
    t.columns[1].width = Cm(23)
    return t

add_section_header(1, "추진배경 및 목적")

add_paragraph("")
p = add_paragraph("", size=10, space_after=4)
add_run_to_paragraph(p, "❍ ", size=10)
add_run_to_paragraph(p, "동일유리(주)는 건축용 복층유리 전문 제조기업으로, 발주서 정리·검토·ERP 입력 및 생산실적 집계 업무에서 ")
add_run_to_paragraph(p, "수작업 의존도가 높아 ", bold=True)
add_run_to_paragraph(p, "업무 효율화가 시급한 상황임")

p = add_paragraph("", size=10, space_after=4)
add_run_to_paragraph(p, "❍ ", size=10)
add_run_to_paragraph(p, "충북대학교 AX기업혁신 공동연구센터와 동일유리 간 산학공동 기술개발과제를 통해 ")
add_run_to_paragraph(p, "AI/자동화(AX) 기반 업무혁신 방안", bold=True)
add_run_to_paragraph(p, "을 도출하고, PBL(Problem-Based Learning) 방식의 현장 중심 교육을 실시하고자 함")

p = add_paragraph("", size=10, space_after=4)
add_run_to_paragraph(p, "❍ ", size=10)
add_run_to_paragraph(p, "동일유리 임직원 대상 현업 데이터 기반 AS-IS 분석 결과를 공유하고, MVP(최소기능제품) 시연을 통해 ")
add_run_to_paragraph(p, "AX 개발 방향에 대한 공감대를 형성", bold=True)
add_run_to_paragraph(p, "하는 것이 본 세미나의 목적임")

add_paragraph("", space_after=6)

# ===================================================
# 2. 프로그램 개요
# ===================================================

add_section_header(2, "프로그램 개요")
add_paragraph("")

items = [
    ("프로그램명", "동일유리 업무자동화(AX) 프로젝트 킥오프 세미나"),
    ("일      시", "2026. 08. 29(토) 10:00 ~ 17:00 (7시간)"),
    ("장      소", "충북대학교 인문사회관 (N14동) 404호"),
    ("대      상", "충북대학교 연구진 4명, 동일유리 직원 20명 (총 24명)"),
    ("행사내용", "AS-IS 업무 분석 발표, AX 개발 방향 논의, MVP 시연, PBL 워크숍"),
]

for label, value in items:
    p = add_paragraph("", size=10, space_after=3)
    add_run_to_paragraph(p, f"❍ {label} : ", size=10, bold=True)
    add_run_to_paragraph(p, value, size=10)

add_paragraph("", space_after=6)

# ===================================================
# 3. 세부 일정
# ===================================================

add_section_header(3, "세부 일정")
add_paragraph("")

schedule = [
    # (구분, 시간, 내용, 비고)
    ("오전", "10:00~10:20 ('20)", "접수 및 등록, 세미나 취지 안내", ""),
    ("", "10:20~11:00 ('40)", "주제발표 ①「동일유리 AX 프로젝트 개요」\n  - 전체 업무 흐름 5단계 소개\n  - AX 대상 2개 영역 설명 (업무지원 / 생산실적)", "충북대 연구진"),
    ("", "11:00~12:00 ('60)", "주제발표 ②「AS-IS 업무 분석 결과」\n  - 발주서 프로세스 4단계 상세 분석\n  - 발주서 6종 양식 비교 및 데이터 변환 과정\n  - 발주서↔작업의뢰서 검증 현황\n  - 복층 생산실적, 재단 실적(절단일보) 현황", "충북대 연구진"),
    ("점심", "12:00~13:00 ('60)", "점심 식사", ""),
    ("오후", "13:00~14:00 ('60)", "주제발표 ③「TO-BE: AX 개발 방향 제안」\n  - 6개 기능 소개 (발주서 파싱, 규격정리, 검증, 복층/재단/주간 실적)\n  - 3단계 개발 전략 (MVP → 확장 → 고도화)\n  - 품명 변환 체계, 위치 정보 구조 분석", "충북대 연구진"),
    ("", "14:00~14:10 ('10)", "휴식", ""),
    ("", "14:10~15:10 ('60)", "MVP 시연 및 토론\n  - 발주서 자동 파싱 데모\n  - 규격 그룹핑 자동화 데모\n  - 발주서↔작업의뢰서 자동 비교 검증 데모\n  - 현업 피드백 수렴", "충북대 연구진\n동일유리 담당자"),
    ("", "15:10~16:10 ('60)", "PBL 워크숍: 현업 과제 도출\n  - 확인 필요사항 Q&A (발주서 양식 통일, ERP 연동 등)\n  - 업무별 Pain Point 심층 논의\n  - 추가 데이터 요청 및 우선순위 결정\n  - 생산현장 입력 방식 논의 (MES 연동 등)", "전체 참석자"),
    ("", "16:10~16:20 ('10)", "휴식", ""),
    ("", "16:20~16:50 ('30)", "향후 계획 수립\n  - 개발 일정 및 마일스톤 합의\n  - 역할 분담 (데이터 제공, 개발, 검증)\n  - 차기 세미나 일정 협의", "전체 참석자"),
    ("", "16:50~17:00 ('10)", "마무리 및 폐회", ""),
]

# 테이블 생성
t = doc.add_table(rows=1 + len(schedule), cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER

# 열 너비 설정
t.columns[0].width = Cm(2.0)
t.columns[1].width = Cm(4.5)
t.columns[2].width = Cm(14.5)
t.columns[3].width = Cm(4.5)

# 헤더
headers = ["구분", "시  간", "세 부  내 용", "비고"]
for i, h in enumerate(headers):
    cell = t.cell(0, i)
    set_cell_shading(cell, "2B2B2B")
    format_cell(cell, h, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

# 데이터 행
for r, (gubun, time_str, content, note) in enumerate(schedule):
    row_idx = r + 1

    # 구분
    cell_g = t.cell(row_idx, 0)
    format_cell(cell_g, gubun, size=9, bold=True if gubun else False)

    # 시간
    cell_t = t.cell(row_idx, 1)
    format_cell(cell_t, time_str, size=9)

    # 내용
    cell_c = t.cell(row_idx, 2)
    cell_c.text = ""
    p = cell_c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)

    lines = content.split("\n")
    for li, line in enumerate(lines):
        if li == 0:
            run = p.add_run(line)
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.name = '맑은 고딕'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        else:
            p2 = cell_c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            run = p2.add_run(line)
            run.font.size = Pt(8.5)
            run.font.name = '맑은 고딕'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # 비고
    cell_n = t.cell(row_idx, 3)
    format_cell(cell_n, note, size=9)

    # 점심/휴식 행 배경색
    if "점심" in content or "휴식" in content or "마무리" in content:
        for ci in range(4):
            set_cell_shading(t.cell(row_idx, ci), "F2F2F2")

# 구분 셀 병합
# 오전: 행1~3 (인덱스 1,2,3)
t.cell(1, 0).merge(t.cell(3, 0))
format_cell(t.cell(1, 0), "오전", size=10, bold=True)
# 점심: 행4 (인덱스 4) - 단독
# 오후: 행5~11 (인덱스 5~11)
t.cell(5, 0).merge(t.cell(11, 0))
format_cell(t.cell(5, 0), "오후", size=10, bold=True)

add_paragraph("", space_after=10)

# ===================================================
# 4. 소요 예산
# ===================================================

add_section_header(4, "소요 예산")
add_paragraph("")

p = add_paragraph("", size=10, space_after=6)
add_run_to_paragraph(p, "❍ 총 소요예산: 약 ", size=10)
add_run_to_paragraph(p, "1,068,000원", size=10, bold=True)

p = add_paragraph("", size=10, space_after=4)
add_run_to_paragraph(p, "❍ 세부 산출내역", size=10, bold=True)

# 예산 테이블
bt = doc.add_table(rows=5, cols=4)
bt.alignment = WD_TABLE_ALIGNMENT.CENTER
bt.columns[0].width = Cm(3.5)
bt.columns[1].width = Cm(12.0)
bt.columns[2].width = Cm(3.5)
bt.columns[3].width = Cm(3.0)

# 헤더
budget_headers = ["구분", "산출내역", "금액 (원)", "비고"]
for i, h in enumerate(budget_headers):
    cell = bt.cell(0, i)
    set_cell_shading(cell, "2B2B2B")
    format_cell(cell, h, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

# 국내여비
format_cell(bt.cell(1, 0), "국내여비", size=10)
bt.cell(1, 0).merge(bt.cell(3, 0))

format_cell(bt.cell(1, 1), "[일비] 25,000원/인·일 x 24인 x 1일 = 600,000원", size=10, alignment=WD_ALIGN_PARAGRAPH.LEFT)
format_cell(bt.cell(2, 1), "[식비] 8,400원/인·일 x 24인 x 1일 = 201,600원", size=10, alignment=WD_ALIGN_PARAGRAPH.LEFT)
format_cell(bt.cell(3, 1), "[다과비] 약 266,400원 (24인 기준)", size=10, alignment=WD_ALIGN_PARAGRAPH.LEFT)

bt.cell(1, 2).merge(bt.cell(3, 2))
format_cell(bt.cell(1, 2), "1,068,000", size=10, bold=True)

bt.cell(1, 3).merge(bt.cell(3, 3))
format_cell(bt.cell(1, 3), "연구진 4인\n+ 동일유리 20인", size=9)

# 합계
set_cell_shading(bt.cell(4, 0), "F2F2F2")
format_cell(bt.cell(4, 0), "합계", size=10, bold=True)
set_cell_shading(bt.cell(4, 1), "F2F2F2")
format_cell(bt.cell(4, 1), "", size=10)
set_cell_shading(bt.cell(4, 2), "F2F2F2")
format_cell(bt.cell(4, 2), "1,068,000", size=10, bold=True)
set_cell_shading(bt.cell(4, 3), "F2F2F2")
format_cell(bt.cell(4, 3), "", size=10)

add_paragraph("", space_after=10)

# ===================================================
# 5. 기대효과
# ===================================================

add_section_header(5, "기대효과")
add_paragraph("")

effects = [
    ("동일유리 현업 담당자와의 직접 소통을 통해 ", "실제 업무 Pain Point를 정밀하게 파악", "하고, 현장에 즉시 적용 가능한 AX 솔루션 개발 방향을 수립할 수 있음"),
    ("AS-IS 분석 결과와 MVP 시연을 통해 동일유리 임직원의 ", "AX에 대한 이해도와 수용성을 높이고", ", 향후 시스템 도입 시 현업 협조 및 데이터 제공의 원활한 기반을 마련함"),
    ("발주서 자동 파싱·검증 및 생산실적 자동 집계 시스템 구축을 통해 ", "수작업 시간 절감, 오류 방지, 생산성 향상", " 등 실질적 업무혁신 효과가 기대됨"),
    ("PBL 방식의 산학협력 교육을 통해 ", "대학의 연구역량과 기업의 현장 경험이 결합", "된 실용적 AX 기술 개발 모델을 확립함"),
]

for normal1, bold_text, normal2 in effects:
    p = add_paragraph("", size=10, space_after=4)
    add_run_to_paragraph(p, "❍ ", size=10)
    add_run_to_paragraph(p, normal1, size=10)
    add_run_to_paragraph(p, bold_text, size=10, bold=True)
    add_run_to_paragraph(p, normal2, size=10)

# ===================================================
# 저장
# ===================================================

output_path = "data/동일유리_세미나_운영_계획(안).docx"
doc.save(output_path)
print(f"세미나 운영 계획(안) 저장 완료: {output_path}")
